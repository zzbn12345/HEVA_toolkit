"""Regression tests that generate their binary extraction inputs at runtime."""

from __future__ import annotations

from pathlib import Path

import docx
import fitz
import pytest
from docx.shared import RGBColor

from heva.extraction.docx_extractor import RunColorIndex, extract_docx_highlights
from heva.extraction.errors import ExtractionCancelled
from heva.extraction.pdf_extractor import (
    PDFTextExtractionError,
    VerticalRectIndex,
    extract_colored_highlights,
    find_highlight_color,
    find_text_span_color,
    group_words_by_block,
    iter_sentence_word_spans,
    require_readable_text_layer,
    sort_page_blocks,
)


def create_highlighted_pdf(path: Path) -> None:
    """Create a minimal PDF with a yellow drawing behind two words."""

    document = fitz.open()
    page = document.new_page()
    text = "Historic harbour remains protected."
    page.insert_text((100, 100), text, fontsize=12)
    highlighted = page.search_for("Historic harbour")[0]
    page.draw_rect(highlighted, fill=(1, 1, 0), color=None, overlay=False)
    document.save(path)
    document.close()


def create_colored_docx(path: Path) -> None:
    """Create a minimal Word document with two custom-colored phrases."""

    document = docx.Document()
    paragraph = document.add_paragraph()
    purple = paragraph.add_run("Historic harbour")
    purple.font.color.rgb = RGBColor(0x70, 0x30, 0xA0)
    paragraph.add_run(" remains ")
    green = paragraph.add_run("protected")
    green.font.color.rgb = RGBColor(0x92, 0xD0, 0x50)
    paragraph.add_run(".")
    document.save(path)


def test_pdf_extractor_reads_generated_highlight_and_mapping(tmp_path: Path) -> None:
    source = tmp_path / "highlighted.pdf"
    create_highlighted_pdf(source)

    raw = extract_colored_highlights(source)
    mapped = extract_colored_highlights(source, color_label_map={"#FFFF00": "historic"})

    assert raw[0]["sentence"] == "Historic harbour remains protected."
    assert raw[0]["entities"] == [
        {
            "start": 0,
            "end": 16,
            "text": "Historic harbour",
            "label": "#FFFF00",
            "color": "#FFFF00",
        }
    ]
    assert mapped[0]["entities"][0]["label"] == "historic"


def test_pdf_extractor_reports_each_source_page_before_sentence_nlp(tmp_path: Path) -> None:
    """Page progress reflects completed adapter reads without claiming NLP completion."""

    source = tmp_path / "two-pages.pdf"
    document = fitz.open()
    for text in ("Historic harbour remains.", "Political use continues."):
        page = document.new_page()
        page.insert_text((100, 100), text, fontsize=12)
        highlighted = page.search_for(text.split()[0])[0]
        page.draw_rect(highlighted, fill=(1, 1, 0), color=None, overlay=False)
    document.save(source)
    document.close()
    progress: list[tuple[int, int]] = []

    records = extract_colored_highlights(
        source,
        progress_callback=lambda completed, total: progress.append((completed, total)),
    )

    assert len(records) == 2
    assert progress == [(1, 2), (2, 2)]


def test_pdf_extractor_cancels_at_page_boundary(tmp_path: Path) -> None:
    """Cancellation stops before the next page and never returns partial records."""

    source = tmp_path / "two-pages.pdf"
    document = fitz.open()
    for text in ("Historic harbour remains.", "Political use continues."):
        page = document.new_page()
        page.insert_text((100, 100), text, fontsize=12)
        highlighted = page.search_for(text.split()[0])[0]
        page.draw_rect(highlighted, fill=(1, 1, 0), color=None, overlay=False)
    document.save(source)
    document.close()
    progress: list[tuple[int, int]] = []

    with pytest.raises(ExtractionCancelled, match="before source page 2"):
        extract_colored_highlights(
            source,
            progress_callback=lambda completed, total: progress.append((completed, total)),
            cancellation_callback=lambda: bool(progress),
        )

    assert progress == [(1, 2)]


def test_pdf_extractor_reads_only_selected_source_pages(tmp_path: Path) -> None:
    source = tmp_path / "three-pages.pdf"
    document = fitz.open()
    for text in ("Historic first.", "Political second.", "Economic third."):
        page = document.new_page()
        page.insert_text((100, 100), text, fontsize=12)
        highlighted = page.search_for(text.split()[0])[0]
        page.draw_rect(highlighted, fill=(1, 1, 0), color=None, overlay=False)
    document.save(source)
    document.close()
    progress: list[tuple[int, int]] = []

    records = extract_colored_highlights(
        source,
        page_numbers=[2, 3],
        progress_callback=lambda completed, total: progress.append((completed, total)),
    )

    assert [record["page"] for record in records] == [2, 3]
    assert all("first" not in record["sentence"] for record in records)
    assert progress == [(1, 2), (2, 2)]


def test_table_pdf_with_broken_font_encoding_fails_before_persisting_gibberish() -> None:
    """The authorized Alpha table fixture must fail honestly instead of yielding cipher text."""

    source = Path(__file__).parents[1] / "2011 EC Galle part 1(Appendix IV) - Copy.pdf"

    with pytest.raises(
        PDFTextExtractionError, match="searchable Unicode text or apply OCR"
    ) as failure:
        extract_colored_highlights(source)

    assert failure.value.code == "pdf_text_unreadable"


def test_text_quality_guard_accepts_short_or_non_latin_language_evidence() -> None:
    """The encoding guard must not mistake short labels or Unicode scripts for corruption."""

    require_readable_text_layer("12345")
    require_readable_text_layer("文化遺産の価値を説明する文章です。" * 30)


def test_pdf_extractor_rejects_a_corrupt_page_inside_otherwise_readable_pdf(
    tmp_path: Path,
) -> None:
    """Readable pages must not dilute a later page's broken font-map evidence."""

    source = tmp_path / "mixed-text-quality.pdf"
    document = fitz.open()
    readable = document.new_page()
    readable_text = "Historic harbour remains important to the community. " * 8
    readable.insert_textbox(fitz.Rect(50, 50, 550, 500), readable_text, fontsize=10)
    highlight = readable.search_for("Historic harbour")[0]
    readable.draw_rect(highlight, fill=(1, 1, 0), color=None, overlay=False)
    corrupted = document.new_page()
    corrupted_text = "!@#$%^&*()_+-=[]{};:',.<>/? " * 12
    corrupted.insert_textbox(fitz.Rect(50, 50, 550, 500), corrupted_text, fontsize=10)
    document.save(source)
    document.close()

    with pytest.raises(PDFTextExtractionError, match="source page 2"):
        extract_colored_highlights(source)


def test_docx_extractor_reads_generated_custom_font_colors(tmp_path: Path) -> None:
    source = tmp_path / "colored.docx"
    create_colored_docx(source)

    actual = extract_docx_highlights(source)

    assert actual[0]["sentence"] == "Historic harbour remains protected."
    assert [entity["text"] for entity in actual[0]["entities"]] == [
        "Historic harbour",
        "protected",
    ]
    assert [entity["color"] for entity in actual[0]["entities"]] == [
        "#7030A0",
        "#92D050",
    ]


def test_docx_extractor_maps_each_generated_color(tmp_path: Path) -> None:
    source = tmp_path / "colored.docx"
    create_colored_docx(source)

    actual = extract_docx_highlights(
        source,
        color_label_map={"#7030A0": "historic", "#92D050": "social"},
    )

    assert actual[0]["values"] == ["historic", "social"]
    assert actual[0]["ner_tags"] == [
        "B-historic",
        "I-historic",
        "O",
        "B-social",
        "O",
    ]


def test_docx_run_color_index_uses_character_overlap_for_dominance() -> None:
    """Tokens crossing run boundaries use the color covering most characters."""
    index = RunColorIndex([
        (0, 3, "#7030A0"),
        (3, 10, "#92D050"),
        (10, 15, None),
    ])

    assert index.dominant_colors([(1, 8), (10, 15)]) == ["#92D050", None]


def test_pdf_sentence_alignment_preserves_order_and_prefix_offsets() -> None:
    """Align ordered word spans without scanning unrelated later sentences."""

    class Sentence:
        def __init__(self, text: str, start: int, end: int) -> None:
            self.text = text
            self.start_char = start
            self.end_char = end

    sentences = [
        Sentence("Label: First sentence.", 0, 22),
        Sentence("Second sentence.", 23, 39),
    ]
    spans = [
        {"word": {"text": "Label:", "color": None}, "start": 0, "end": 6, "page": 1},
        {"word": {"text": "First", "color": "#FFFF00"}, "start": 7, "end": 12, "page": 1},
        {"word": {"text": "sentence.", "color": None}, "start": 13, "end": 22, "page": 1},
        {"word": {"text": "Second", "color": None}, "start": 23, "end": 29, "page": 2},
        {"word": {"text": "sentence.", "color": None}, "start": 30, "end": 39, "page": 2},
    ]

    aligned = list(iter_sentence_word_spans(sentences, spans))

    assert aligned[0][1] == "First sentence."
    assert [item["word"]["text"] for item in aligned[0][3]] == ["First", "sentence."]
    assert [(item["start"], item["end"]) for item in aligned[0][3]] == [(0, 5), (6, 15)]
    assert [item["word"]["text"] for item in aligned[1][3]] == ["Second", "sentence."]
    assert aligned[1][4] == [2, 2]


def test_pdf_block_sort_preserves_band_and_column_reading_order() -> None:
    """Read a heading, then each column, then the following full-width block."""
    blocks = [
        (20.0, 10.0, 580.0, 30.0, "Heading", 0, 0),
        (20.0, 40.0, 180.0, 50.0, "Left one", 1, 0),
        (220.0, 40.0, 380.0, 50.0, "Right one", 2, 0),
        (20.0, 60.0, 180.0, 70.0, "Left two", 3, 0),
        (220.0, 60.0, 380.0, 70.0, "Right two", 4, 0),
        (20.0, 100.0, 580.0, 120.0, "Footer heading", 5, 0),
    ]

    ordered, metadata = sort_page_blocks(
        blocks, fitz.Rect(0.0, 0.0, 600.0, 200.0), rotation=0
    )

    assert [block[5] for block in ordered] == [0, 1, 3, 2, 4, 5]
    assert metadata[1][1] == metadata[3][1] == 0
    assert metadata[2][1] == metadata[4][1] == 1


def test_pdf_highlight_index_preserves_overlap_and_drawing_order() -> None:
    """Spatial candidates must retain the original maximum-overlap behavior."""
    drawings = [
        {"rect": fitz.Rect(0, 100, 10, 110), "color": "#RED000"},
        {"rect": fitz.Rect(0, 0, 10, 10), "color": "#FIRST0"},
        {"rect": fitz.Rect(0, 0, 10, 10), "color": "#SECOND"},
    ]
    word = fitz.Rect(0, 0, 10, 10)
    index = VerticalRectIndex(drawings, bucket_height=20.0)

    assert find_highlight_color(word, drawings) == "#FIRST0"
    assert find_highlight_color(word, drawings, index) == "#FIRST0"


def test_pdf_span_index_preserves_first_containing_span_behavior() -> None:
    """A neutral first containing span must still suppress later colors."""
    spans = [
        {"rect": fitz.Rect(0, 100, 10, 110), "color": "#7030A0"},
        {"rect": fitz.Rect(0, 0, 10, 10), "color": "#000000"},
        {"rect": fitz.Rect(0, 0, 10, 10), "color": "#7030A0"},
    ]
    word = fitz.Rect(0, 0, 10, 10)
    index = VerticalRectIndex(spans, bucket_height=20.0)

    assert find_text_span_color(word, spans) is None
    assert find_text_span_color(word, spans, index) is None


def test_pdf_words_group_by_block_without_changing_order() -> None:
    """A single grouping pass must preserve words as extracted from the page."""
    words = [
        {"text": "left", "block_no": 2},
        {"text": "heading", "block_no": 1},
        {"text": "column", "block_no": 2},
    ]

    grouped = group_words_by_block(words)

    assert [word["text"] for word in grouped[2]] == ["left", "column"]
    assert [word["text"] for word in grouped[1]] == ["heading"]
