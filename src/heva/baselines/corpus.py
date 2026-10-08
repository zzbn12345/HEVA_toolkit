"""Load baseline training data from HEVA Data Packages, structured spreadsheets and
Atlas.ti CSV exports.

Expected layout under the data directory (every part optional)::

    <data_dir>/heva/**/heva-annotations.json   (or heva-annotations.csv)
    <data_dir>/structured/*.xlsx | *.csv        (52-column HEVA spreadsheet format)
    <data_dir>/atlasti/*.csv

HEVA packages feed T1 (sentence values) and T2 (BIO spans). Structured spreadsheets and
Atlas.ti quotations only feed T1: they carry values for whole texts, not token spans.
Every example records its collection, so splits and scores can be made per collection.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from .labels import VALUES, normalize_value

HEVA_JSON = "heva-annotations.json"
HEVA_CSV = "heva-annotations.csv"

# Atlas.ti codes look like "Value: Economic - Before"; the Before/After suffix is dropped.
VALUE_CODE = re.compile(r"Value:\s*([A-Za-z]+)(?:\s*-\s*[A-Za-z]+)?")
# Column order of the Atlas.ti Quotation Manager export, used when the header row was
# saved in a localized UI language and can no longer be decoded (e.g. "????").
QUOTATION_MANAGER_POSITIONS = {"id": 0, "document": 2, "content": 4, "codes": 6}

# Structured spreadsheets (see EXAMPLE.xlsx): one text per row, L1 values in columns
# "1".."8" in VALUES order. A cell counts as present when it is greater than 0, because
# some files hold 2 for a value coded twice (e.g. Atlas.ti "Before" and "After").
STRUCTURED_TEXT = "Sentence"
STRUCTURED_COLLECTION = "Folder_ID"
STRUCTURED_DOCUMENT = "Datasource"
STRUCTURED_VALUE_COLUMNS = tuple(str(code) for code in range(1, len(VALUES) + 1))


@dataclass(frozen=True)
class T1Example:
    """One text with the subset of CVF values it expresses (possibly empty)."""

    example_id: str
    text: str
    values: tuple[str, ...]
    source: str
    document: str
    collection: str = ""


@dataclass(frozen=True)
class T2Example:
    """One tokenized sentence with one BIO tag per token."""

    example_id: str
    tokens: tuple[str, ...]
    tags: tuple[str, ...]
    document: str
    collection: str = ""

    @property
    def values(self) -> tuple[str, ...]:
        return tuple(sorted({tag[2:] for tag in self.tags if tag != "O"}))


@dataclass
class LoadReport:
    """What was read, skipped or merged, so a run can be audited."""

    files: list[str] = field(default_factory=list)
    skipped: dict[str, int] = field(default_factory=dict)
    merged_duplicates: int = 0
    conflicting_duplicates: int = 0
    unknown_value_codes: dict[str, int] = field(default_factory=dict)

    def skip(self, reason: str) -> None:
        self.skipped[reason] = self.skipped.get(reason, 0) + 1

    def as_dict(self) -> dict:
        return {
            "files": self.files,
            "skipped": self.skipped,
            "merged_duplicates": self.merged_duplicates,
            "conflicting_duplicates": self.conflicting_duplicates,
            "unknown_value_codes": self.unknown_value_codes,
        }


def _normalize_text(text: str) -> str:
    return " ".join(text.split())


def _stable_id(prefix: str, text: str) -> str:
    return f"{prefix}:{hashlib.sha1(text.encode('utf-8')).hexdigest()[:12]}"


# --- HEVA Data Packages ---------------------------------------------------------------


def find_heva_package_files(root: Path) -> list[Path]:
    """Return one annotation file per package directory, preferring canonical JSON."""
    if not root.is_dir():
        return []
    directories = {path.parent for path in root.rglob(HEVA_JSON)}
    directories |= {path.parent for path in root.rglob(HEVA_CSV)}
    files = []
    for directory in sorted(directories):
        json_path = directory / HEVA_JSON
        files.append(json_path if json_path.exists() else directory / HEVA_CSV)
    return files


def _read_heva_json(path: Path) -> list[dict]:
    package = json.loads(path.read_text(encoding="utf-8"))
    records = []
    for document in package.get("documents", []):
        for record in document.get("records", []):
            records.append({**record, "document_id": document.get("document_id", path.parent.name)})
    return records


def _read_heva_csv(path: Path) -> list[dict]:
    records = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            record = dict(row)
            for column in ("values", "tokens", "entities", "ner_tags"):
                record[column] = json.loads(row[column]) if row.get(column) else []
            records.append(record)
    return records


def _canonical_tags(tags, report: LoadReport) -> tuple[str, ...] | None:
    canonical = []
    for tag in tags:
        if tag == "O":
            canonical.append(tag)
            continue
        prefix, _, name = str(tag).partition("-")
        value = normalize_value(name)
        if prefix not in ("B", "I") or value is None:
            report.skip("unknown_bio_tag")
            return None
        canonical.append(f"{prefix}-{value}")
    return tuple(canonical)


def load_heva_records(root: Path, report: LoadReport) -> tuple[list[T1Example], list[T2Example]]:
    """Read every HEVA package under ``root`` into T1 and T2 examples."""
    t1, t2 = [], []
    seen_ids = set()
    for path in find_heva_package_files(root):
        report.files.append(str(path))
        # The package folder directly under heva/ names the collection.
        collection = path.parent.relative_to(root).parts[0] if path.parent != root else root.name
        records = _read_heva_json(path) if path.suffix == ".json" else _read_heva_csv(path)
        for record in records:
            example_id = f"heva:{record.get('document_id')}:{record.get('sentence_id')}"
            if example_id in seen_ids:
                report.skip("duplicate_record_id")
                continue
            seen_ids.add(example_id)

            tokens = record.get("tokens") or []
            tags = _canonical_tags(record.get("ner_tags") or [], report)
            if tags is None:
                continue
            if not tokens or len(tokens) != len(tags):
                report.skip("tokens_tags_length_mismatch")
                continue

            values = {normalize_value(str(value)) for value in record.get("values") or []}
            values |= {tag[2:] for tag in tags if tag != "O"}
            if None in values:
                report.skip("unknown_value")
                continue

            document = str(record.get("document_id"))
            text = (record.get("curated_sentence") or record.get("sentence") or "").strip()
            if text:
                t1.append(
                    T1Example(example_id, text, tuple(sorted(values)), "heva", document, collection)
                )
            t2.append(T2Example(example_id, tuple(tokens), tags, document, collection))
    return t1, t2


# --- Atlas.ti CSV exports ---------------------------------------------------------------


def _decode(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252")


def _atlasti_columns(header: list[str], path: Path) -> dict[str, int]:
    names = {name.strip().lower(): index for index, name in enumerate(header)}
    if "quotation content" in names and "codes" in names:
        return {
            "id": names.get("id", 0),
            "document": names.get("document", names.get("document name", 0)),
            "content": names["quotation content"],
            "codes": names["codes"],
        }
    unreadable = all(set(name.strip()) <= {"?"} for name in header[1:] if name.strip())
    if unreadable and len(header) > max(QUOTATION_MANAGER_POSITIONS.values()):
        return dict(QUOTATION_MANAGER_POSITIONS)
    raise ValueError(
        f"{path}: cannot find the 'Quotation Content' and 'Codes' columns of an Atlas.ti export."
    )


def read_atlasti_csv(path: Path, report: LoadReport) -> list[T1Example]:
    """Read one Atlas.ti quotation export; labels come from the 'Value: X' codes."""
    text = _decode(path.read_bytes())
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";"
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        return []
    columns = _atlasti_columns(rows[0], path)

    examples, value_codes_found = [], 0
    for row in rows[1:]:
        if len(row) <= max(columns.values()):
            report.skip("atlasti_short_row")
            continue
        content = _normalize_text(row[columns["content"]])
        if not content:
            report.skip("atlasti_empty_quotation")
            continue
        values = set()
        for code in VALUE_CODE.findall(row[columns["codes"]]):
            value_codes_found += 1
            value = normalize_value(code)
            if value is None:
                report.unknown_value_codes[code] = report.unknown_value_codes.get(code, 0) + 1
            else:
                values.add(value)
        examples.append(
            T1Example(
                _stable_id("atlasti", content),
                content,
                tuple(sorted(values)),
                "atlasti",
                row[columns["document"]].strip(),
                path.stem,
            )
        )
    if rows[1:] and value_codes_found == 0:
        raise ValueError(f"{path}: no 'Value: <name>' codes found; is this an Atlas.ti export?")
    return examples


def load_atlasti_records(root: Path, report: LoadReport) -> list[T1Example]:
    """Read every Atlas.ti CSV under ``root``."""
    if not root.is_dir():
        return []
    examples = []
    for path in sorted(root.rglob("*.csv")):
        report.files.append(str(path))
        examples.extend(read_atlasti_csv(path, report))
    return examples


# --- Structured spreadsheets -----------------------------------------------------------


def _cell_text(cell) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        cell = int(cell)
    return str(cell).strip()


def _is_structured_header(header: list[str]) -> bool:
    return STRUCTURED_TEXT in header and all(code in header for code in STRUCTURED_VALUE_COLUMNS)


def _structured_rows_xlsx(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ImportError(
            f"{path}: reading .xlsx needs openpyxl (pip install -e '.[baselines]'), "
            "or save the sheet as CSV."
        ) from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            if rows and _is_structured_header(rows[0]):
                return rows
    finally:
        workbook.close()
    raise ValueError(f"{path}: no sheet has a '{STRUCTURED_TEXT}' column and value columns 1-8.")


def _structured_rows_csv(path: Path) -> list[list[str]]:
    text = _decode(path.read_bytes())
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ","
    rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    if not rows or not _is_structured_header(rows[0]):
        raise ValueError(f"{path}: no '{STRUCTURED_TEXT}' column and value columns 1-8.")
    return rows


def _present(cell: str, report: LoadReport) -> bool:
    if not cell:
        return False
    try:
        return float(cell) > 0
    except ValueError:
        report.skip("structured_non_numeric_value_cell")
        return False


def read_structured_file(path: Path, report: LoadReport) -> list[T1Example]:
    """Read one spreadsheet in the 52-column HEVA format (one text per row)."""
    rows = _structured_rows_xlsx(path) if path.suffix.lower() == ".xlsx" else _structured_rows_csv(path)
    header = rows[0]
    column = {name: index for index, name in reversed(list(enumerate(header)))}
    value_columns = [column[code] for code in STRUCTURED_VALUE_COLUMNS]

    def cell(row: list[str], name: str) -> str:
        index = column.get(name)
        return row[index] if index is not None and index < len(row) else ""

    examples = []
    for row in rows[1:]:
        if not any(row):
            continue
        text = _normalize_text(cell(row, STRUCTURED_TEXT))
        if not text:
            report.skip("structured_empty_sentence")
            continue
        values = tuple(
            value
            for value, index in zip(VALUES, value_columns)
            if index < len(row) and _present(row[index], report)
        )
        examples.append(
            T1Example(
                _stable_id("structured", text),
                text,
                values,
                "structured",
                cell(row, STRUCTURED_DOCUMENT),
                cell(row, STRUCTURED_COLLECTION) or path.stem,
            )
        )
    return examples


def load_structured_records(root: Path, report: LoadReport) -> list[T1Example]:
    """Read every structured spreadsheet (.xlsx or .csv) under ``root``."""
    if not root.is_dir():
        return []
    examples = []
    for path in sorted(root.rglob("*")):
        if path.suffix.lower() not in (".xlsx", ".csv") or path.name.startswith(("~$", ".")):
            continue
        report.files.append(str(path))
        examples.extend(read_structured_file(path, report))
    return examples


# --- Task datasets ----------------------------------------------------------------------


def deduplicate_t1(examples: list[T1Example], report: LoadReport) -> list[T1Example]:
    """Merge identical texts so they cannot land in both train and test.

    Different exports of one Atlas.ti project repeat the same quotations; their value
    sets are unioned.
    """
    merged: dict[str, T1Example] = {}
    for example in examples:
        key = _normalize_text(example.text)
        if key not in merged:
            merged[key] = example
            continue
        report.merged_duplicates += 1
        kept = merged[key]
        if set(kept.values) != set(example.values):
            report.conflicting_duplicates += 1
            union = tuple(sorted(set(kept.values) | set(example.values)))
            merged[key] = T1Example(
                kept.example_id, kept.text, union, kept.source, kept.document, kept.collection
            )
    return list(merged.values())


def deduplicate_t2(examples: list[T2Example], report: LoadReport) -> list[T2Example]:
    """Keep the first of identical token sequences; tag conflicts cannot be merged."""
    kept: dict[tuple[str, ...], T2Example] = {}
    for example in examples:
        if example.tokens not in kept:
            kept[example.tokens] = example
            continue
        report.merged_duplicates += 1
        if kept[example.tokens].tags != example.tags:
            report.conflicting_duplicates += 1
    return list(kept.values())


def load_t1_dataset(data_dir: Path) -> tuple[list[T1Example], LoadReport]:
    report = LoadReport()
    heva_t1, _ = load_heva_records(data_dir / "heva", report)
    structured = load_structured_records(data_dir / "structured", report)
    atlasti = load_atlasti_records(data_dir / "atlasti", report)
    return deduplicate_t1(heva_t1 + structured + atlasti, report), report


def load_t2_dataset(data_dir: Path) -> tuple[list[T2Example], LoadReport]:
    report = LoadReport()
    _, heva_t2 = load_heva_records(data_dir / "heva", report)
    return deduplicate_t2(heva_t2, report), report
